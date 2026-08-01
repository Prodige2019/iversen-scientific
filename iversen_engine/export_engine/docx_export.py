"""
Export Word (.docx) d'une correction — même contenu que l'export PDF, même moteur
de rendu de formules (matplotlib mathtext), mise en page adaptée au format Word.

Choix technique : `python-docx` plutôt que docx-js (Node). Ce module fait partie
d'un service backend Python redistribuable (appelé depuis l'API FastAPI comme
l'export PDF) — pas une génération ponctuelle de document dans un environnement
interactif. Rester 100% Python évite d'ajouter une dépendance Node à l'ensemble
du projet pour cette seule fonctionnalité.
"""
import os
import tempfile
from typing import Any, Dict

from docx import Document
from docx.shared import Pt, Mm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .latex_render import render_latex_to_png

INK = RGBColor(0x1B, 0x2A, 0x45)
INK_SOFT = RGBColor(0x4B, 0x5A, 0x73)
RED = RGBColor(0x8A, 0x3B, 0x2F)
GOLD = RGBColor(0x8A, 0x6A, 0x15)
GREEN = RGBColor(0x3E, 0x7A, 0x52)


def _add_formula_image(doc: Document, latex_str: str, tmpdir: str, idx: int, max_width_in: float = 5.5):
    path = os.path.join(tmpdir, f"formula_{idx}.png")
    render_latex_to_png(latex_str, path)
    from PIL import Image as PILImage
    with PILImage.open(path) as im:
        w_px, h_px = im.size
    width_in = min(max_width_in, w_px / 200)  # dpi=200 utilisé par render_latex_to_png
    height_in = width_in * (h_px / w_px)
    doc.add_picture(path, width=Inches(width_in), height=Inches(height_in))


def export_correction_to_docx(correction: Dict[str, Any], output_path: str) -> None:
    """Génère un .docx à partir du dict produit par engine.correction_to_dict()."""
    doc = Document()

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Mm(20)

    with tempfile.TemporaryDirectory() as tmpdir:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(correction.get("exercise_title", "Correction"))
        title_run.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = INK

        score_p = doc.add_paragraph()
        score_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        score_run = score_p.add_run(f"{correction.get('score', '')}/{correction.get('score_max', 20)}")
        score_run.bold = True
        score_run.font.size = Pt(14)
        score_run.font.color.rgb = GOLD

        subject_line = (correction.get("subject_line") or f"f(x) = {correction.get('function_str', '')}").replace("$", "")
        subject_p = doc.add_paragraph()
        subject_run = subject_p.add_run(subject_line)
        subject_run.italic = True
        subject_run.font.size = Pt(11)
        subject_run.font.color.rgb = INK_SOFT

        doc.add_paragraph()

        img_idx = 0
        for i, step in enumerate(correction.get("steps", []), start=1):
            step_title = doc.add_paragraph()
            step_title.paragraph_format.space_before = Pt(10)
            run = step_title.add_run(f"Étape {i} — {step['title']}")
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = INK

            if step.get("result_latex"):
                img_idx += 1
                try:
                    _add_formula_image(doc, step["result_latex"], tmpdir, img_idx)
                except Exception:
                    p = doc.add_paragraph()
                    r = p.add_run(step["result_latex"])
                    r.font.color.rgb = INK

            if step.get("explanation"):
                p = doc.add_paragraph()
                r = p.add_run(step["explanation"])
                r.font.size = Pt(10.5)
                r.font.color.rgb = INK_SOFT

            if step.get("rule_recalled"):
                img_idx += 1
                try:
                    label_p = doc.add_paragraph()
                    label_run = label_p.add_run("Règle rappelée :")
                    label_run.font.size = Pt(9)
                    label_run.font.color.rgb = GREEN
                    label_run.italic = True
                    _add_formula_image(doc, step["rule_recalled"], tmpdir, img_idx, max_width_in=5.0)
                except Exception:
                    pass

            if step.get("warning"):
                warn_p = doc.add_paragraph()
                warn_run = warn_p.add_run("⚠ " + step["warning"])
                warn_run.font.size = Pt(9.5)
                warn_run.font.color.rgb = RED

        doc.add_paragraph()
        footer_p = doc.add_paragraph()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer_p.add_run(
            "Iversen Scientific — correction générée automatiquement (moteur symbolique, sans IA générative)"
        )
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = INK_SOFT

        doc.save(output_path)
